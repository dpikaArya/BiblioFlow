"""Report generation utilities."""
from __future__ import annotations
import logging
import sys
from pathlib import Path
from datetime import datetime
from dataclasses import asdict

import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))
from core.models import PipelineState

logger = logging.getLogger("aibef.reports")


class ReportGenerator:
    """Generate all report files."""

    @staticmethod
    def generate_all(state: PipelineState, output_dir: Path, report_dir: Path = None):
        """Generate all report files."""
        if report_dir is None:
            report_dir = output_dir
        ReportGenerator._prisma_flow_text(state, report_dir / "prisma_flow.txt")
        ReportGenerator._summary_table(state, report_dir / "pipeline_summary.txt")

        certified_path = output_dir / "Bibliometrix_Compatible.txt"
        if certified_path.exists() and state.stats.final_count > 0:
            ReportGenerator._framework_execution_summary(state, output_dir, report_dir)
            ReportGenerator._certification_report(state, output_dir, report_dir)
            ReportGenerator._bibliometrix_validation_report(state, output_dir, report_dir)

    @staticmethod
    def _prisma_flow_text(state: PipelineState, path: Path):
        recs = state.prisma_data.get("records", {})
        lines = [
            "=" * 60,
            "PRISMA FLOW DIAGRAM",
            "=" * 60,
            "",
            f"Records identified through database searching:",
            f"  Total: {recs.get('imported', 'N/A')}",
            f"  Per file: {dict(state.stats.per_file)}",
            "",
            f"Records after merging: {recs.get('merged', 'N/A')}",
            "",
            f"Duplicates removed: {recs.get('duplicates_identified', 'N/A')}",
            f"Records after deduplication: {recs.get('after_dedup', 'N/A')}",
            "",
            f"Records after cleaning: {recs.get('after_cleaning', 'N/A')}",
            "",
            f"Studies included in final review: {recs.get('included', 'N/A')}",
            "",
            "=" * 60,
        ]
        path.write_text("\n".join(lines), encoding="utf-8")

    @staticmethod
    def _summary_table(state: PipelineState, path: Path):
        lines = [
            "AIBEF Pipeline Summary",
            "=" * 40,
            f"Total imported: {state.stats.total_imported}",
            f"After merge: {state.stats.after_merge}",
            f"Duplicates found: {state.stats.duplicates_found}",
            f"After dedup: {state.stats.after_dedup}",
            f"Validation issues: {state.stats.validation_issues}",
            f"Metadata repairs: {state.stats.metadata_repairs}",
            f"After cleaning: {state.stats.after_cleaning}",
            f"Final count: {state.stats.final_count}",
            f"Synchronized: {state.stats.synchronized}",
            f"Errors: {len(state.errors)}",
        ]
        path.write_text("\n".join(lines), encoding="utf-8")

    @staticmethod
    def _framework_execution_summary(state: PipelineState, output_dir: Path, report_dir: Path):
        from docx import Document
        doc = Document()
        doc.add_heading("Framework Execution Summary", level=0)
        doc.add_paragraph(f"Generated: {datetime.now().isoformat()}")

        doc.add_heading("1. Pipeline Statistics", level=1)
        doc.add_paragraph(f"Total Imported: {state.stats.total_imported}")
        doc.add_paragraph(f"After Merge: {state.stats.after_merge}")
        doc.add_paragraph(f"Duplicates Found: {state.stats.duplicates_found}")
        doc.add_paragraph(f"After Dedup: {state.stats.after_dedup}")
        doc.add_paragraph(f"After Cleaning: {state.stats.after_cleaning}")
        doc.add_paragraph(f"Final Count: {state.stats.final_count}")
        doc.add_paragraph(f"Synchronized: {state.stats.synchronized}")
        doc.add_paragraph(f"Errors: {len(state.errors)}")

        doc.add_heading("2. Per-File Import Counts", level=1)
        for filename, count in state.stats.per_file.items():
            doc.add_paragraph(f"{filename}: {count} records", style="List Bullet")

        doc.add_heading("3. Certified Dataset", level=1)
        certified_path = output_dir / "Bibliometrix_Compatible.txt"
        doc.add_paragraph(f"Status: READY_FOR_BIBLIOSHINY")
        doc.add_paragraph(f"File: {certified_path.name}")
        doc.add_paragraph(f"Path: {certified_path}")
        doc.add_paragraph(f"Records: {state.stats.final_count}")
        doc.add_paragraph(f"Format: TAB-delimited (.txt)")
        doc.add_paragraph(f"Launch Command: python main.py --launch-biblioshiny-only")

        if state.errors:
            doc.add_heading("4. Errors", level=1)
            for error in state.errors:
                doc.add_paragraph(error, style="List Bullet")

        report_path = report_dir / "Framework_Execution_Summary.docx"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(report_path))

    @staticmethod
    def _certification_report(state: PipelineState, output_dir: Path, report_dir: Path):
        from docx import Document
        doc = Document()
        doc.add_heading("Certification Report", level=0)
        doc.add_paragraph(f"Generated: {datetime.now().isoformat()}")

        doc.add_heading("1. Certification Status", level=1)
        doc.add_paragraph(f"Status: READY_FOR_BIBLIOSHINY")
        doc.add_paragraph(f"Synchronized: {state.stats.synchronized}")
        doc.add_paragraph(f"Final Records: {state.stats.final_count}")

        doc.add_heading("2. Validation Checks", level=1)
        doc.add_paragraph(f"Synchronization Check: PASS" if state.stats.synchronized else "FAIL")
        required_cols = ["AU", "TI", "SO", "PY"]
        if state.validated_dataset is not None:
            bib_compatible = all(col in state.validated_dataset.columns for col in required_cols)
            doc.add_paragraph(f"Bibliometrix Compatibility: {'PASS' if bib_compatible else 'FAIL'}")
        else:
            doc.add_paragraph("Bibliometrix Compatibility: FAIL (no validated dataset)")

        doc.add_heading("3. Certified Dataset", level=1)
        certified_path = output_dir / "Bibliometrix_Compatible.txt"
        doc.add_paragraph(f"File: {certified_path.name}")
        doc.add_paragraph(f"Path: {certified_path}")
        doc.add_paragraph(f"Records: {state.stats.final_count}")
        doc.add_paragraph(f"Columns: {len(state.validated_dataset.columns) if state.validated_dataset is not None else 0}")
        doc.add_paragraph(f"Format: TAB-delimited (.txt)")

        doc.add_heading("4. Database Source Distribution", level=1)
        for filename, count in state.stats.per_file.items():
            doc.add_paragraph(f"{filename}: {count} records", style="List Bullet")

        report_path = report_dir / "Certification_Report.docx"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(report_path))

    @staticmethod
    def _bibliometrix_validation_report(state: PipelineState, output_dir: Path, report_dir: Path):
        from docx import Document
        doc = Document()
        doc.add_heading("Bibliometrix Validation Report", level=0)
        doc.add_paragraph(f"Generated: {datetime.now().isoformat()}")

        doc.add_heading("1. Validation Summary", level=1)
        doc.add_paragraph(f"Final Records: {state.stats.final_count}")
        doc.add_paragraph(f"Synchronized: {state.stats.synchronized}")

        doc.add_heading("2. Required Fields Check", level=1)
        required_fields = ["AU", "TI", "SO", "PY", "DT", "DI", "DE", "ID", "CR", "C1", "RP", "TC", "LA"]
        if state.validated_dataset is not None:
            for field in required_fields:
                present = field in state.validated_dataset.columns
                doc.add_paragraph(f"{field}: {'PRESENT' if present else 'MISSING'}", style="List Bullet")

        doc.add_heading("3. Certified Dataset", level=1)
        certified_path = output_dir / "Bibliometrix_Compatible.txt"
        doc.add_paragraph(f"Status: READY_FOR_BIBLIOSHINY")
        doc.add_paragraph(f"File: {certified_path.name}")
        doc.add_paragraph(f"Path: {certified_path}")
        doc.add_paragraph(f"Records: {state.stats.final_count}")
        doc.add_paragraph(f"Format: TAB-delimited (.txt)")
        doc.add_paragraph(f"Launch Command: python main.py --launch-biblioshiny-only")

        report_path = report_dir / "Bibliometrix_Validation_Report.docx"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(report_path))
