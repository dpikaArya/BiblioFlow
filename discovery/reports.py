"""Report generators for Multi-Database Discovery Engine."""
from __future__ import annotations
import json
import logging
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))

from discovery.models import (
    DatabaseGroup,
    DiscoveryReport,
    FileRecord,
    GroupPipelineResult,
    SchemaValidationResult,
)

logger = logging.getLogger("aibef.discovery.reports")


class DiscoveryReportGenerator:
    """Generates all discovery-related reports."""

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_identification_report(
        self, file_records: list[FileRecord]
    ) -> Path:
        """Generate Database_Identification_Report.xlsx."""
        path = self.output_dir / "Database_Identification_Report.xlsx"
        rows = []
        for r in file_records:
            rows.append({
                "Filename": r.file_name,
                "Detected Database": r.database_source.value,
                "Detected Format": r.export_format.value,
                "Confidence": f"{r.confidence:.2%}",
                "Confidence Level": r.confidence_level.value,
                "Detection Method": r.detection_method,
                "Schema Version": r.schema_version or "N/A",
                "Encoding": r.encoding,
                "Delimiter": r.delimiter or "N/A",
                "Record Count": r.record_count,
                "Status": r.status.value if hasattr(r.status, "value") else str(r.status),
                "Errors": "; ".join(r.errors) if r.errors else "None",
                "Warnings": "; ".join(r.warnings) if r.warnings else "None",
            })

        df = pd.DataFrame(rows)
        df.to_excel(path, index=False, engine="openpyxl")
        logger.info("Generated identification report: %s", path)
        return path

    def generate_identification_warning_report(
        self, warnings: list
    ) -> Optional[Path]:
        """Generate Database_Identification_Warning.docx for low-confidence files."""
        if not warnings:
            logger.info("No identification warnings to report")
            return None

        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            logger.warning("python-docx not available, generating text warning report")
            return self._generate_warning_text(warnings)

        path = self.output_dir / "Database_Identification_Warning.docx"
        doc = Document()
        title = doc.add_heading("Database Identification Warning Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        doc.add_paragraph("")

        for w in warnings:
            doc.add_heading(f"File: {w.file_name}", level=1)

            table = doc.add_table(rows=5, cols=2, style="Table Grid")
            fields = [
                ("Reason", w.reason),
                ("Evidence", w.evidence),
                ("Suggested Database", w.suggested_database),
                ("Suggested Repair", w.suggested_repair),
                ("User Action Required", w.user_action_required),
            ]
            for i, (label, value) in enumerate(fields):
                table.cell(i, 0).text = label
                table.cell(i, 1).text = value
            doc.add_paragraph("")

        doc.save(str(path))
        logger.info("Generated warning report: %s", path)
        return path

    def _generate_warning_text(self, warnings: list) -> Path:
        path = self.output_dir / "Database_Identification_Warning.txt"
        lines = [
            "DATABASE IDENTIFICATION WARNING REPORT",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "=" * 60,
            "",
        ]
        for w in warnings:
            lines.extend([
                f"File: {w.file_name}",
                f"  Reason: {w.reason}",
                f"  Evidence: {w.evidence}",
                f"  Suggested Database: {w.suggested_database}",
                f"  Suggested Repair: {w.suggested_repair}",
                f"  User Action Required: {w.user_action_required}",
                "",
            ])
        path.write_text("\n".join(lines), encoding="utf-8")
        logger.info("Generated text warning report: %s", path)
        return path

    def generate_schema_validation_report(
        self, schema_results: list[SchemaValidationResult]
    ) -> Path:
        """Generate Schema_Validation_Report.docx."""
        path = self.output_dir / "Schema_Validation_Report.docx"
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return self._generate_schema_text_report(schema_results, path)

        doc = Document()
        title = doc.add_heading("Schema Validation Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        doc.add_paragraph("")

        summary_table = doc.add_table(rows=4, cols=2, style="Table Grid")
        total = len(schema_results)
        passed = sum(1 for sr in schema_results if sr.is_valid)
        failed = total - passed
        summary_table.cell(0, 0).text = "Total Files Validated"
        summary_table.cell(0, 1).text = str(total)
        summary_table.cell(1, 0).text = "Passed"
        summary_table.cell(1, 1).text = str(passed)
        summary_table.cell(2, 0).text = "Failed"
        summary_table.cell(2, 1).text = str(failed)
        summary_table.cell(3, 0).text = "Pass Rate"
        summary_table.cell(3, 1).text = f"{(passed/total*100):.1f}%" if total > 0 else "N/A"
        doc.add_paragraph("")

        for sr in schema_results:
            doc.add_heading(f"File: {sr.file_record.file_name}", level=1)
            status_text = "PASSED" if sr.is_valid else "FAILED"
            doc.add_paragraph(f"Database: {sr.file_record.database_source.value}")
            doc.add_paragraph(f"Format: {sr.file_record.export_format.value}")
            doc.add_paragraph(f"Validation Score: {sr.validation_score:.1f}/100")
            doc.add_paragraph(f"Status: {status_text}")
            doc.add_paragraph(
                f"Mandatory Fields: {sr.mandatory_fields_present}/{sr.mandatory_fields_total}"
            )
            doc.add_paragraph(
                f"Recommended Fields: {sr.recommended_fields_present}/{sr.recommended_fields_total}"
            )

            checks = [
                ("Encoding", sr.encoding_valid),
                ("Delimiter", sr.delimiter_valid),
                ("Column Types", sr.column_types_valid),
                ("Identifier Integrity", sr.identifier_integrity_valid),
                ("Author Structure", sr.author_structure_valid),
                ("Keyword Structure", sr.keyword_structure_valid),
                ("Record Integrity", sr.record_integrity_valid),
            ]
            doc.add_heading("Validation Checks", level=2)
            check_table = doc.add_table(rows=len(checks) + 1, cols=2, style="Table Grid")
            check_table.cell(0, 0).text = "Check"
            check_table.cell(0, 1).text = "Status"
            for i, (name, passed_check) in enumerate(checks):
                check_table.cell(i + 1, 0).text = name
                check_table.cell(i + 1, 1).text = "PASSED" if passed_check else "FAILED"

            if sr.issues:
                doc.add_heading("Issues", level=2)
                for issue in sr.issues:
                    doc.add_paragraph(issue, style="List Bullet")
            if sr.warnings:
                doc.add_heading("Warnings", level=2)
                for warning in sr.warnings:
                    doc.add_paragraph(warning, style="List Bullet")

        doc.save(str(path))
        logger.info("Generated schema validation report: %s", path)
        return path

    def _generate_schema_text_report(
        self, results: list[SchemaValidationResult], path: Path
    ) -> Path:
        lines = [
            "SCHEMA VALIDATION REPORT",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "=" * 60,
            "",
        ]
        for sr in results:
            lines.extend([
                f"File: {sr.file_record.file_name}",
                f"  Database: {sr.file_record.database_source.value}",
                f"  Score: {sr.validation_score:.1f}/100",
                f"  Valid: {sr.is_valid}",
                f"  Mandatory: {sr.mandatory_fields_present}/{sr.mandatory_fields_total}",
                f"  Issues: {sr.issues}",
                f"  Warnings: {sr.warnings}",
                "",
            ])
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    def generate_framework_execution_summary(
        self, report: DiscoveryReport
    ) -> Path:
        """Generate Framework_Execution_Summary.docx."""
        path = self.output_dir / "Framework_Execution_Summary.docx"
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return self._generate_summary_text(report, path)

        doc = Document()
        title = doc.add_heading("Framework Execution Summary", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        doc.add_paragraph(f"Run ID: {report.run_id}")
        doc.add_paragraph(f"Search Directory: {report.search_directory}")
        doc.add_paragraph("")

        doc.add_heading("1. Discovery Summary", level=1)
        t = doc.add_table(rows=6, cols=2, style="Table Grid")
        summary_data = [
            ("Files Discovered", str(report.files_discovered)),
            ("Files Supported", str(report.files_supported)),
            ("Files Unsupported", str(report.files_unsupported)),
            ("Database Groups", str(report.database_groups)),
            ("Total Records", str(report.total_records)),
            ("Discovery Timestamp", report.discovery_timestamp),
        ]
        for i, (label, val) in enumerate(summary_data):
            t.cell(i, 0).text = label
            t.cell(i, 1).text = val
        doc.add_paragraph("")

        if report.groups:
            doc.add_heading("2. Database Groups", level=1)
            gt = doc.add_table(rows=len(report.groups) + 1, cols=5, style="Table Grid")
            gt.cell(0, 0).text = "Group"
            gt.cell(0, 1).text = "Database"
            gt.cell(0, 2).text = "Files"
            gt.cell(0, 3).text = "Records"
            gt.cell(0, 4).text = "Pipeline Status"
            for i, g in enumerate(report.groups):
                gt.cell(i + 1, 0).text = g.group_name
                gt.cell(i + 1, 1).text = g.database_source.value
                gt.cell(i + 1, 2).text = str(len(g.files))
                gt.cell(i + 1, 3).text = str(g.total_records)
                gt.cell(i + 1, 4).text = g.pipeline_status.value
            doc.add_paragraph("")

        if report.pipeline_results:
            doc.add_heading("3. Pipeline Results", level=1)
            for pr in report.pipeline_results:
                doc.add_heading(f"{pr.group_name}", level=2)
                pt = doc.add_table(rows=7, cols=2, style="Table Grid")
                pr_data = [
                    ("Status", pr.status.value),
                    ("Records Imported", str(pr.records_imported)),
                    ("Records Final", str(pr.records_final)),
                    ("Quality Score", f"{pr.quality_score}/100 ({pr.quality_grade})"),
                    ("Certification", pr.certification_status),
                    ("Duration", f"{pr.duration_seconds:.1f}s"),
                    ("Errors", str(len(pr.errors))),
                ]
                for i, (label, val) in enumerate(pr_data):
                    pt.cell(i, 0).text = label
                    pt.cell(i, 1).text = val
                doc.add_paragraph("")

        doc.add_heading("4. Cross-Database Merge Status", level=1)
        doc.add_paragraph(f"Status: {report.cross_database_merge_status}")
        doc.add_paragraph(
            "Cross-database merging is BLOCKED. Each database group "
            "is processed independently."
        )
        doc.add_paragraph("")

        doc.add_heading("5. Validation Status", level=1)
        statuses = [
            ("Cross-Database Merge", report.cross_database_merge_status),
            ("Framework Health Score", f"{report.framework_health_score:.1f}/100"),
            ("Release Readiness", report.release_readiness),
            ("Overall Status", report.overall_status.value),
        ]
        st = doc.add_table(rows=len(statuses), cols=2, style="Table Grid")
        for i, (label, val) in enumerate(statuses):
            st.cell(i, 0).text = label
            st.cell(i, 1).text = val

        if report.errors:
            doc.add_heading("6. Errors", level=1)
            for error in report.errors:
                doc.add_paragraph(error, style="List Bullet")

        if report.warnings:
            doc.add_heading("7. Warnings", level=1)
            for warning in report.warnings:
                doc.add_paragraph(warning, style="List Bullet")

        doc.save(str(path))
        logger.info("Generated framework execution summary: %s", path)
        return path

    def _generate_summary_text(
        self, report: DiscoveryReport, path: Path
    ) -> Path:
        lines = [
            "FRAMEWORK EXECUTION SUMMARY",
            f"Run ID: {report.run_id}",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "=" * 60,
            "",
            f"Files Discovered: {report.files_discovered}",
            f"Files Supported: {report.files_supported}",
            f"Files Unsupported: {report.files_unsupported}",
            f"Database Groups: {report.database_groups}",
            f"Total Records: {report.total_records}",
            f"Cross-Database Merge: {report.cross_database_merge_status}",
            f"Health Score: {report.framework_health_score:.1f}/100",
            f"Release Readiness: {report.release_readiness}",
            f"Overall Status: {report.overall_status.value}",
            "",
        ]
        for g in report.groups:
            lines.append(
                f"  {g.group_name}: {g.database_source.value} "
                f"({len(g.files)} files, {g.total_records} records)"
            )
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    def generate_provenance_log(
        self, report: DiscoveryReport, path: Optional[Path] = None
    ) -> Path:
        """Generate Provenance_Log.json for the discovery run."""
        if path is None:
            path = self.output_dir / "Provenance_Log.json"

        data = {
            "framework": "AIBEF",
            "version": "2.0.0",
            "module": "MultiDatabaseDiscoveryEngine",
            "run_id": report.run_id,
            "generated": datetime.now().isoformat(),
            "search_directory": report.search_directory,
            "files_discovered": report.files_discovered,
            "database_groups": report.database_groups,
            "total_records": report.total_records,
            "groups": [],
        }

        for g in report.groups:
            group_data = {
                "group_id": g.group_id,
                "group_name": g.group_name,
                "database_source": g.database_source.value,
                "files": [
                    {
                        "file_name": f.file_name,
                        "file_path": str(f.file_path),
                        "database_source": f.database_source.value,
                        "export_format": f.export_format.value,
                        "confidence": f.confidence,
                        "record_count": f.record_count,
                        "encoding": f.encoding,
                        "discovery_timestamp": f.discovery_timestamp,
                    }
                    for f in g.files
                ],
            }
            data["groups"].append(group_data)

        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False, default=str)
        logger.info("Generated provenance log: %s", path)
        return path

    def generate_audit_log(
        self, report: DiscoveryReport, path: Optional[Path] = None
    ) -> Path:
        """Generate Audit_Log.json."""
        if path is None:
            path = self.output_dir / "Audit_Log.json"

        data = {
            "framework": "AIBEF",
            "version": "2.0.0",
            "module": "MultiDatabaseDiscoveryEngine",
            "run_id": report.run_id,
            "generated": datetime.now().isoformat(),
            "events": [],
            "errors": report.errors,
            "warnings": report.warnings,
        }

        for fr in report.file_records:
            data["events"].append({
                "event": "file_identified",
                "file_name": fr.file_name,
                "database_source": fr.database_source.value,
                "confidence": fr.confidence,
                "status": fr.status.value if hasattr(fr.status, "value") else str(fr.status),
            })

        for g in report.groups:
            data["events"].append({
                "event": "group_created",
                "group_name": g.group_name,
                "database_source": g.database_source.value,
                "file_count": len(g.files),
                "total_records": g.total_records,
            })

        for pr in report.pipeline_results:
            data["events"].append({
                "event": "pipeline_executed",
                "group_name": pr.group_name,
                "status": pr.status.value,
                "records_final": pr.records_final,
                "duration_seconds": pr.duration_seconds,
            })

        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False, default=str)
        logger.info("Generated audit log: %s", path)
        return path
