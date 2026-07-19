"""Biblioshiny Launch Manager - Core launcher."""
from __future__ import annotations
import hashlib
import json
import os
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document

from .config import BiblioshinyConfig
from .models import LaunchPrechecks, LaunchResult, DatasetInfo, REnvironmentInfo, ServerDetectionResult
from .validator import check_certification, verify_r_environment, locate_certified_dataset, run_full_prechecks
from .r_interface import launch_biblioshiny
from .server_discovery import ShinyServerDiscovery
from .browser import open_validated_browser


class BiblioshinyLaunchManager:
    """Manages the launch of Biblioshiny from AIBEF certified output."""

    def __init__(self, config: Optional[BiblioshinyConfig] = None):
        self.config = config or BiblioshinyConfig()

    def check_certification(self, output_dir: Path) -> list:
        return check_certification(output_dir)

    def verify_r_environment(self) -> REnvironmentInfo:
        return verify_r_environment(self.config)

    def verify_bibliometrix(self) -> bool:
        info = verify_r_environment(self.config)
        return info.bibliometrix_installed and info.load_test_passed

    def locate_certified_dataset(self, output_dir: Path) -> DatasetInfo:
        return locate_certified_dataset(output_dir, self.config)

    def launch_biblioshiny(self, dataset_path: Path) -> LaunchResult:
        result, _proc = launch_biblioshiny(dataset_path, self.config)
        result.browser_status = "launched" if result.success else result.browser_status
        return result

    def generate_launch_report(
        self,
        prechecks: LaunchPrechecks,
        launch_result: Optional[LaunchResult],
        output_path: Path,
    ) -> Path:
        """Generate Biblioshiny_Launch_Report.docx."""
        doc = Document()
        doc.add_heading("Biblioshiny Launch Report", level=0)
        doc.add_paragraph(f"Generated: {datetime.now().isoformat()}")

        doc.add_heading("1. Dataset Information", level=1)
        if prechecks.dataset and prechecks.dataset.exists:
            doc.add_paragraph(f"Path: {prechecks.dataset.path}")
            doc.add_paragraph(f"Records: {prechecks.dataset.record_count}")
            doc.add_paragraph(f"Fields: {prechecks.dataset.field_count}")
            doc.add_paragraph(f"Hash: {prechecks.dataset.file_hash}")
            doc.add_paragraph(f"Certification ID: {prechecks.dataset.certification_id}")
        else:
            doc.add_paragraph("No certified dataset found.")

        doc.add_heading("2. Certification Evidence", level=1)
        for check in prechecks.certification_checks:
            status = "PASS" if check.passed else "FAIL"
            doc.add_paragraph(f"[{status}] {check.name}: {check.message}")
        doc.add_paragraph(f"Overall: {prechecks.certification_status}")

        doc.add_heading("3. R Environment", level=1)
        if prechecks.r_environment:
            r = prechecks.r_environment
            doc.add_paragraph(f"R Version: {r.r_version}")
            doc.add_paragraph(f"Bibliometrix Installed: {r.bibliometrix_installed}")
            doc.add_paragraph(f"Bibliometrix Version: {r.bibliometrix_version}")
            doc.add_paragraph(f"Load Test Passed: {r.load_test_passed}")
            if r.errors:
                doc.add_paragraph("Errors:")
                for err in r.errors:
                    doc.add_paragraph(f"  - {err}", style="List Bullet")
        else:
            doc.add_paragraph("R environment not checked.")

        doc.add_heading("4. Launch Status", level=1)
        if launch_result:
            doc.add_paragraph(f"Success: {launch_result.success}")
            doc.add_paragraph(f"Process ID: {launch_result.process_id}")
            if launch_result.error_message:
                doc.add_paragraph(f"Error: {launch_result.error_message}")

            detection = launch_result.server_detection
            if detection and detection.detected:
                doc.add_heading("4a. Server Detection", level=2)
                doc.add_paragraph(f"Detection Method: {detection.detection_method}")
                doc.add_paragraph(f"Detected URL: {detection.url}")
                doc.add_paragraph(f"Host: {detection.host}")
                doc.add_paragraph(f"Port: {detection.port}")
                doc.add_paragraph(f"HTTP Status: {detection.http_status}")
                doc.add_paragraph(f"Server Reachable: {detection.server_reachable}")
                doc.add_paragraph(f"Shiny Application: {detection.shiny_application_detected}")
                doc.add_paragraph(f"Biblioshiny Identified: {detection.biblioshiny_identified}")
                doc.add_paragraph(f"Response Latency: {detection.response_latency_ms:.1f}ms")
                doc.add_paragraph(f"Detection Time: {detection.detection_time_seconds:.1f}s")
                if detection.probe_attempts:
                    doc.add_paragraph(f"Probe Attempts: {detection.probe_attempts}")
                doc.add_heading("Browser Launch", level=3)
                doc.add_paragraph(f"Status: {detection.browser_status}")
                if detection.manual_url:
                    doc.add_paragraph(f"Manual URL: {detection.manual_url}")
            elif detection and detection.error_message:
                doc.add_heading("4a. Server Detection", level=2)
                doc.add_paragraph(f"Detection Method: {detection.detection_method}")
                doc.add_paragraph(f"Error: {detection.error_message}")
                if detection.manual_url:
                    doc.add_paragraph(f"Manual URL: {detection.manual_url}")
        else:
            doc.add_paragraph("Launch not attempted.")

        doc.add_heading("5. User Instructions", level=1)
        doc.add_paragraph(
            "Biblioshiny provides an interactive Shiny-based interface for bibliometric analysis. "
            "Upon opening, use the File menu to import your certified dataset, or it may auto-load "
            "depending on configuration."
        )
        doc.add_paragraph("To import the certified dataset manually:")
        doc.add_paragraph("1. Click 'File' > 'Import'", style="List Number")
        doc.add_paragraph(f"2. Navigate to: {prechecks.dataset.path if prechecks.dataset else 'N/A'}", style="List Number")
        doc.add_paragraph("3. Select 'Web of Science' as the database source", style="List Number")
        doc.add_paragraph("4. Click 'Start Analysis'", style="List Number")

        if launch_result and launch_result.server_detection:
            det = launch_result.server_detection
            if det.manual_url:
                doc.add_heading("6. Troubleshooting", level=1)
                doc.add_paragraph("If the browser did not open automatically:")
                doc.add_paragraph(f"Open this URL manually: {det.manual_url}")
                doc.add_paragraph(
                    "Ensure the R process is still running. "
                    "If the server is not responding, restart Biblioshiny "
                    "using the AIBEF launch command."
                )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def generate_launch_log(
        self,
        prechecks: LaunchPrechecks,
        launch_result: Optional[LaunchResult],
        output_path: Path,
    ) -> Path:
        """Generate Biblioshiny_Launch_Log.json."""
        log_data = {
            "run_id": f"blm_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
            "dataset_path": prechecks.dataset.path if prechecks.dataset else "",
            "dataset_hash": prechecks.dataset.file_hash if prechecks.dataset else "",
            "certification_status": prechecks.certification_status,
            "r_version": prechecks.r_environment.r_version if prechecks.r_environment else "",
            "bibliometrix_version": prechecks.r_environment.bibliometrix_version if prechecks.r_environment else "",
            "launch_time": launch_result.launch_time if launch_result else "",
            "process_id": launch_result.process_id if launch_result else None,
            "success": launch_result.success if launch_result else False,
            "error": launch_result.error_message if launch_result else "",
            "timestamp": datetime.now().isoformat(),
        }

        if launch_result and launch_result.server_detection:
            det = launch_result.server_detection
            log_data["server_detection"] = {
                "detection_method": det.detection_method,
                "server_url": det.url,
                "host": det.host,
                "port": det.port,
                "http_status": det.http_status,
                "detection_time_seconds": det.detection_time_seconds,
                "probe_attempts": det.probe_attempts,
                "server_verification": {
                    "reachable": det.server_reachable,
                    "shiny_application": det.shiny_application_detected,
                    "biblioshiny_identified": det.biblioshiny_identified,
                    "response_latency_ms": det.response_latency_ms,
                },
                "browser_launch_status": launch_result.browser_status,
                "manual_url": det.manual_url,
                "console_detection_succeeded": det.console_detection_succeeded,
                "fallback_detection_succeeded": det.fallback_detection_succeeded,
            }
            if det.probe_results:
                log_data["server_detection"]["probe_results"] = [
                    p.to_dict() for p in det.probe_results
                ]
        else:
            log_data["server_detection"] = None

        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(log_data, indent=2, ensure_ascii=False), encoding="utf-8")
        return output_path

    def generate_block_report(
        self,
        prechecks: LaunchPrechecks,
        output_path: Path,
    ) -> Path:
        """Generate Biblioshiny_Launch_Block_Report.docx when launch is blocked."""
        doc = Document()
        doc.add_heading("Biblioshiny Launch Block Report", level=0)
        doc.add_paragraph(f"Generated: {datetime.now().isoformat()}")

        doc.add_heading("Certification Status", level=1)
        doc.add_paragraph(f"Status: {prechecks.certification_status}")

        doc.add_heading("Failure Reason", level=1)
        doc.add_paragraph(prechecks.block_reason or "Unknown")

        doc.add_heading("Certification Checks", level=1)
        for check in prechecks.certification_checks:
            status = "PASS" if check.passed else "FAIL"
            doc.add_paragraph(f"[{status}] {check.name}: {check.message}")

        if prechecks.r_environment and prechecks.r_environment.errors:
            doc.add_heading("R Environment Issues", level=1)
            for err in prechecks.r_environment.errors:
                doc.add_paragraph(f"  - {err}", style="List Bullet")

        doc.add_heading("Recommended Action", level=1)
        if "Certification" in prechecks.block_reason:
            doc.add_paragraph(
                "Run the full AIBEF certification pipeline and ensure all stages pass "
                "before attempting to launch Biblioshiny."
            )
            doc.add_paragraph("Command: python main.py")
        elif "R environment" in prechecks.block_reason or "R_NOT_READY" in prechecks.certification_status:
            doc.add_paragraph(
                "Ensure R >= 4.3 is installed and the bibliometrix package is available. "
                "Run: Rscript -e \"install.packages('bibliometrix')\""
            )
        elif "Dataset" in prechecks.block_reason:
            doc.add_paragraph(
                "The certified Bibliometrix_Compatible.txt was not found in the output directory. "
                "Run the pipeline first."
            )
        else:
            doc.add_paragraph("Resolve the issues listed above and retry.")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def run(
        self,
        output_dir: Path,
        report_dir: Optional[Path] = None,
        launch: bool = True,
    ) -> dict:
        """Execute the full Biblioshiny launch workflow.

        Returns a dict with prechecks, launch_result, and generated file paths.
        """
        if report_dir is None:
            report_dir = output_dir

        prechecks = run_full_prechecks(output_dir, self.config)

        result = {
            "prechecks": prechecks.to_dict(),
            "launch_result": None,
            "files_generated": [],
        }

        if not prechecks.all_passed and not self.config.allow_uncertified_launch:
            block_path = report_dir / "Biblioshiny_Launch_Block_Report.docx"
            self.generate_block_report(prechecks, block_path)
            result["files_generated"].append(str(block_path))
            return result

        launch_result = None
        if launch and prechecks.dataset and prechecks.dataset.exists:
            ds_path = Path(prechecks.dataset.path)

            self._print_launch_summary(prechecks)

            launch_result, proc = launch_biblioshiny(ds_path, self.config)

            if launch_result.success:
                discovery = ShinyServerDiscovery(self.config.server_detection)
                detection_result = discovery.discover(
                    process=proc,
                    timeout_seconds=self.config.launch_timeout_seconds,
                )
                launch_result.server_detection = detection_result

                if detection_result.detected and self.config.auto_open_browser:
                    detection_result = open_validated_browser(
                        detection_result,
                        auto_open=self.config.auto_open_browser,
                    )
                    launch_result.browser_status = detection_result.browser_status
                elif detection_result.url:
                    launch_result.browser_status = "manual_url_ready"
                    detection_result.manual_url = detection_result.url
                else:
                    launch_result.browser_status = "detection_failed"

            result["launch_result"] = launch_result.to_dict()

            if proc is not None:
                try:
                    proc.terminate()
                except Exception:
                    pass

        report_path = report_dir / "Biblioshiny_Launch_Report.docx"
        self.generate_launch_report(prechecks, launch_result, report_path)
        result["files_generated"].append(str(report_path))

        log_path = report_dir / "Biblioshiny_Launch_Log.json"
        self.generate_launch_log(prechecks, launch_result, log_path)
        result["files_generated"].append(str(log_path))

        return result

    def _print_launch_summary(self, prechecks: LaunchPrechecks):
        """Print the pre-launch summary to console."""
        ds = prechecks.dataset
        r = prechecks.r_environment

        print()
        print("=" * 60)
        print("AIBEF Biblioshiny Launch")
        print("=" * 60)
        print(f"Certification:   {prechecks.certification_status}")
        print(f"Dataset:         {Path(ds.path).name if ds else 'N/A'}")
        print(f"Records:         {ds.record_count if ds else 'N/A'}")
        print(f"Fields:          {ds.field_count if ds else 'N/A'}")
        print(f"Dataset Hash:    {ds.file_hash if ds else 'N/A'}")
        print(f"R Version:       {r.r_version if r else 'N/A'}")
        print(f"Bibliometrix:    {r.bibliometrix_version if r else 'N/A'}")
        print(f"Detection Mode:  {self.config.server_detection.mode}")
        for check in prechecks.certification_checks:
            status = "PASS" if check.passed else "FAIL"
            print(f"  [{status}] {check.name}")
        print("-" * 60)
        print("Launching R process...")
        print("=" * 60)
